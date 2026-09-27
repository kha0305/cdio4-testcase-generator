"""
Router: /api/projects

Quan ly toan dien cac Du An (Projects):
- Danh sach du an kem KPI chi so kiem thu (Requirements, Test Suites, Pass/Fail)
- Tao du an moi (Chon Ca Nhan hoac Nhom, tu dong tao ma tham gia join_code)
- Tham gia du an bang ma (POST /api/projects/join)
- Xem chi tiet du an (cac requirement, parameters, suites)
- Chinh sua du an (Name, Code, Version, Lead, Project Type)
- Xoa du an (Cascade delete)
- Xuat bao cao kiem thu toan dien cap Du An (Excel 4 sheet, CSV, JSON, Markdown)
"""

import io
import json
import secrets
import datetime
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from database import get_db
from models import Project, Requirement, Parameter, TestSuite, TestCase, ProjectMember, User
from schemas import (
    ProjectCreate, ProjectUpdate, ProjectOut, ProjectJoinRequest,
    ProjectMemberOut, ProjectMemberAdd, ProjectMemberRoleUpdate,
)
from routers.auth import get_auth_user

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

router = APIRouter(prefix="/api/projects", tags=["Quản lý Dự án"])

ROLE_NAMES = {
    "leader": "Trưởng nhóm",
    "deputy": "Phó nhóm",
    "member": "Thành viên",
}


def _get_user_role(project_id: int, user_id: int, db: Session) -> str:
    """Xac dinh vai tro cua nguoi dung trong du an (leader, deputy, member)."""
    mem = db.query(ProjectMember).filter(
        ProjectMember.project_id == project_id,
        ProjectMember.user_id == user_id,
    ).first()
    if mem:
        return mem.role
    
    project = db.query(Project).filter(Project.id == project_id).first()
    if project and project.owner_id == user_id:
        return "leader"
    return "member"


def _calculate_project_stats(project: Project, db: Session = None, current_user_id: int = None):
    """Tinh toan chi so thong ke KPI va thong tin thanh vien cho 1 du an."""
    if db is not None:
        req_ids = [r[0] for r in db.query(Requirement.id).filter(Requirement.project_id == project.id).all()]
        req_count = len(req_ids)
        if req_ids:
            suite_ids = [s[0] for s in db.query(TestSuite.id).filter(TestSuite.requirement_id.in_(req_ids)).all()]
            suite_count = len(suite_ids)
            if suite_ids:
                cases = db.query(TestCase.status).filter(TestCase.test_suite_id.in_(suite_ids)).all()
                total_cases = len(cases)
                pass_cnt = sum(1 for (st,) in cases if st and st.strip().lower() in ("pass", "passed", "đạt"))
                fail_cnt = sum(1 for (st,) in cases if st and st.strip().lower() in ("fail", "failed", "không đạt", "lỗi"))
                untested_cnt = total_cases - pass_cnt - fail_cnt
                pass_rate = round((pass_cnt / total_cases * 100), 1) if total_cases > 0 else 0.0
            else:
                total_cases = pass_cnt = fail_cnt = untested_cnt = 0
                pass_rate = 0.0
        else:
            suite_count = total_cases = pass_cnt = fail_cnt = untested_cnt = 0
            pass_rate = 0.0

        # Dem so thanh vien
        members_count = db.query(ProjectMember).filter(ProjectMember.project_id == project.id).count()
        if members_count == 0:
            members_count = 1

        # Vai tro cua current_user trong du an
        cur_role = "member"
        if current_user_id:
            cur_role = _get_user_role(project.id, current_user_id, db)
    else:
        req_count = len(project.requirements)
        suites = []
        for r in project.requirements:
            suites.extend(r.test_suites)
        suite_count = len(suites)
        all_cases = []
        for s in suites:
            all_cases.extend(s.test_cases)
        total_cases = len(all_cases)
        pass_cnt = sum(1 for tc in all_cases if tc.status and tc.status.strip().lower() in ("pass", "passed", "đạt"))
        fail_cnt = sum(1 for tc in all_cases if tc.status and tc.status.strip().lower() in ("fail", "failed", "không đạt", "lỗi"))
        untested_cnt = total_cases - pass_cnt - fail_cnt
        pass_rate = round((pass_cnt / total_cases * 100), 1) if total_cases > 0 else 0.0
        members_count = len(project.members) if hasattr(project, "members") and project.members else 1
        cur_role = "leader" if (current_user_id and getattr(project, "owner_id", None) == current_user_id) else "member"

    created_str = project.created_at.strftime("%d/%m/%Y") if project.created_at else ""
    owner_name = project.lead or "QA Lead"
    if project.owner:
        owner_name = project.owner.full_name

    return {
        "id": project.id,
        "name": project.name,
        "code": project.code or f"PRJ-{project.id:03d}",
        "description": project.description or "",
        "version": project.version or "1.0.0",
        "lead": owner_name,
        "project_type": getattr(project, "project_type", "personal") or "personal",
        "join_code": getattr(project, "join_code", "") or "",
        "owner_id": getattr(project, "owner_id", None),
        "owner_name": owner_name,
        "current_user_role": cur_role,
        "members_count": max(members_count, 1),
        "created_at": created_str,
        "requirements_count": req_count,
        "test_suites_count": suite_count,
        "total_cases": total_cases,
        "pass_count": pass_cnt,
        "fail_count": fail_cnt,
        "untested_count": untested_cnt,
        "pass_rate": pass_rate,
    }


@router.get(
    "",
    response_model=list[ProjectOut],
    summary="Danh sách toàn bộ dự án",
    description="Lấy danh sách tất cả các dự án (Cá nhân & Nhóm) kèm chỉ số thống kê KPI kiểm thử và vai trò của bạn trong dự án.",
)
def get_all_projects(current_user: User = Depends(get_auth_user), db: Session = Depends(get_db)):
    """Lấy danh sách tất cả dự án kèm thống kê KPI và vai trò người dùng."""
    projects = db.query(Project).order_by(Project.id.desc()).all()
    
    # Nếu chưa có dự án nào, tự động tạo dự án mẫu khởi điểm
    if not projects:
        default_p = Project(
            name="Dự Án Đăng Ký Tài Khoản & Phân Quyền",
            code="PRJ-AUTH",
            description="Hệ thống xác thực người dùng, phân quyền vai trò và kiểm tra dữ liệu đầu vào.",
            version="1.0.0",
            lead=current_user.full_name,
            owner_id=current_user.id,
            project_type="personal",
        )
        db.add(default_p)
        db.commit()
        db.refresh(default_p)
        # Thêm current_user làm leader
        mem = ProjectMember(project_id=default_p.id, user_id=current_user.id, role="leader")
        db.add(mem)
        db.commit()
        projects = [default_p]

    return [_calculate_project_stats(p, db, current_user.id) for p in projects]


@router.post(
    "",
    response_model=ProjectOut,
    summary="Tạo dự án mới (Cá nhân hoặc Nhóm)",
    description="Khởi tạo một dự án kiểm thử mới. Người tạo mặc định là Trưởng nhóm (Leader) và có quyền phân quyền cho thành viên khác.",
)
def create_project(payload: ProjectCreate, current_user: User = Depends(get_auth_user), db: Session = Depends(get_db)):
    """Tạo dự án mới - Người tạo mặc định là Leader."""
    proj_type = payload.project_type or "personal"
    join_code = None
    if proj_type == "team":
        join_code = payload.join_code or f"TEAM-{secrets.token_hex(3).upper()}"

    p = Project(
        name=payload.name,
        code=payload.code or f"PRJ-{datetime.datetime.utcnow().strftime('%M%S')}",
        description=payload.description or "",
        version=payload.version or "1.0.0",
        lead=current_user.full_name,
        owner_id=current_user.id,
        project_type=proj_type,
        join_code=join_code,
    )
    db.add(p)
    db.commit()
    db.refresh(p)

    # Tự động gán người tạo làm Leader duy nhất khởi điểm
    mem = ProjectMember(project_id=p.id, user_id=current_user.id, role="leader")
    db.add(mem)
    db.commit()

    return _calculate_project_stats(p, db, current_user.id)


@router.post(
    "/join",
    response_model=ProjectOut,
    summary="Tham gia dự án nhóm bằng mã mời",
    description="Thành viên tham gia vào dự án nhóm bằng mã mời tham gia (Join Code) hoặc mã dự án. Mặc định vai trò là Thành viên (Member).",
)
def join_project_by_code(payload: ProjectJoinRequest, current_user: User = Depends(get_auth_user), db: Session = Depends(get_db)):
    """Tham gia vào một dự án nhóm bằng mã join_code hoặc mã code."""
    query_code = payload.join_code.strip().upper()
    project = db.query(Project).filter(
        (Project.join_code == query_code) | (Project.code == query_code)
    ).first()
    if not project:
        raise HTTPException(
            status_code=404,
            detail=f"Không tìm thấy dự án với mã '{query_code}'. Vui lòng kiểm tra lại mã tham gia."
        )

    # Thêm thành viên vào ProjectMember nếu chưa có
    existing_mem = db.query(ProjectMember).filter(
        ProjectMember.project_id == project.id,
        ProjectMember.user_id == current_user.id,
    ).first()
    if not existing_mem:
        # Nếu là người tạo thì role là leader, ngược lại là member
        new_role = "leader" if project.owner_id == current_user.id else "member"
        new_mem = ProjectMember(project_id=project.id, user_id=current_user.id, role=new_role)
        db.add(new_mem)
        db.commit()

    return _calculate_project_stats(project, db, current_user.id)


@router.get(
    "/{project_id}",
    summary="Chi tiết dự án và các chức năng con",
    description="Lấy thông tin chi tiết của dự án bao gồm danh sách yêu cầu, bộ tham số bóc tách và các bộ Test Suite.",
)
def get_project_details(project_id: int, current_user: User = Depends(get_auth_user), db: Session = Depends(get_db)):
    """Xem chi tiết 1 dự án kèm danh sách requirement và test suites."""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Không tìm thấy dự án")

    stats = _calculate_project_stats(project, db, current_user.id)

    # Danh sach requirements cua du an
    reqs_out = []
    for r in project.requirements:
        suites_out = []
        for s in r.test_suites:
            suites_out.append({
                "id": s.id,
                "technique": s.technique,
                "total_cases": s.total_cases,
                "created_at": s.created_at.strftime("%d/%m/%Y %H:%M") if s.created_at else "",
                "coverage_info": s.coverage_info,
            })

        reqs_out.append({
            "id": r.id,
            "title": r.title,
            "raw_text": r.raw_text,
            "format_type": r.format_type,
            "assignee_name": r.assignee_name or "",
            "due_date": r.due_date or "",
            "priority": r.priority or "Medium",
            "status": r.status or "planning",
            "created_at": r.created_at.strftime("%d/%m/%Y %H:%M") if r.created_at else "",
            "parameters": [
                {
                    "id": p.id,
                    "name": p.name,
                    "data_type": p.data_type,
                    "min_value": p.min_value,
                    "max_value": p.max_value,
                    "min_length": p.min_length,
                    "max_length": p.max_length,
                    "enum_values": p.enum_values,
                }
                for p in r.parameters
            ],
            "test_suites": suites_out,
        })

    return {
        **stats,
        "project": stats,
        "requirements": reqs_out,
    }


@router.put(
    "/{project_id}",
    response_model=ProjectOut,
    summary="Cập nhật thông tin dự án",
    description="Chỉnh sửa tên dự án, mã code, phiên bản, người phụ trách và chế độ dự án. Chỉ Trưởng nhóm (Leader) hoặc Phó nhóm (Deputy) mới có quyền thực hiện.",
)
def update_project(project_id: int, payload: ProjectUpdate, current_user: User = Depends(get_auth_user), db: Session = Depends(get_db)):
    """Chỉnh sửa thông tin dự án (Chỉ Leader hoặc Deputy)."""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Không tìm thấy dự án")

    user_role = _get_user_role(project_id, current_user.id, db)
    if user_role not in ("leader", "deputy") and project.owner_id != current_user.id:
        raise HTTPException(
            status_code=403,
            detail="Bạn không có quyền chỉnh sửa dự án này. Chỉ Trưởng nhóm (Leader) hoặc Phó nhóm mới có quyền."
        )

    if payload.name is not None:
        project.name = payload.name
    if payload.code is not None:
        project.code = payload.code
    if payload.description is not None:
        project.description = payload.description
    if payload.version is not None:
        project.version = payload.version
    if payload.lead is not None:
        project.lead = payload.lead
    if payload.project_type is not None:
        project.project_type = payload.project_type
        if project.project_type == "team" and not project.join_code:
            project.join_code = f"TEAM-{secrets.token_hex(3).upper()}"

    db.commit()
    db.refresh(project)
    return _calculate_project_stats(project, db, current_user.id)


@router.delete(
    "/{project_id}",
    summary="Xóa dự án và toàn bộ dữ liệu",
    description="Xóa vĩnh viễn dự án cùng toàn bộ dữ liệu. Chỉ duy nhất Trưởng nhóm (Leader) người sở hữu mới có quyền xóa.",
)
def delete_project(project_id: int, current_user: User = Depends(get_auth_user), db: Session = Depends(get_db)):
    """Xóa dự án và toàn bộ requirements, test cases liên quan (Chỉ Leader)."""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Không tìm thấy dự án")

    user_role = _get_user_role(project_id, current_user.id, db)
    if user_role != "leader" and project.owner_id != current_user.id:
        raise HTTPException(
            status_code=403,
            detail="Chỉ duy nhất Trưởng nhóm (Leader) mới có quyền xóa vĩnh viễn dự án này."
        )

    db.delete(project)
    db.commit()
    return {"message": f"Dự án {project_id} đã được xóa thành công"}


# ---------------------------------------------------------------------------
# QUẢN LÝ THÀNH VIÊN & PHÂN QUYỀN TRONG DỰ ÁN (PROJECT RBAC)
# ---------------------------------------------------------------------------
@router.get(
    "/{project_id}/members",
    summary="Danh sách thành viên và phân quyền trong dự án",
    description="Lấy danh sách toàn bộ thành viên trong dự án kèm vai trò (Leader, Deputy, Member) và vai trò của người dùng hiện tại.",
)
def get_project_members(project_id: int, current_user: User = Depends(get_auth_user), db: Session = Depends(get_db)):
    """Lấy danh sách thành viên dự án."""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Không tìm thấy dự án")

    current_role = _get_user_role(project_id, current_user.id, db)
    members = db.query(ProjectMember).filter(ProjectMember.project_id == project_id).all()

    # Neu chua co member nao (du lieu cu), tao ban ghi cho owner
    if not members and project.owner_id:
        owner_mem = ProjectMember(project_id=project_id, user_id=project.owner_id, role="leader")
        db.add(owner_mem)
        db.commit()
        members = [owner_mem]

    results = []
    for m in members:
        u = m.user or db.query(User).filter(User.id == m.user_id).first()
        if u:
            results.append({
                "id": m.id,
                "project_id": m.project_id,
                "user_id": u.id,
                "username": u.username,
                "email": u.email,
                "full_name": u.full_name,
                "role": m.role,
                "role_name": ROLE_NAMES.get(m.role, "Thành viên"),
                "joined_at": m.joined_at.strftime("%d/%m/%Y") if m.joined_at else "",
            })

    return {
        "project_id": project_id,
        "project_name": project.name,
        "current_user_id": current_user.id,
        "current_user_role": current_role,
        "current_user_role_name": ROLE_NAMES.get(current_role, "Thành viên"),
        "members": results,
    }


@router.post(
    "/{project_id}/members",
    summary="Thêm thành viên vào dự án",
    description="Trưởng nhóm hoặc Phó nhóm thêm một thành viên vào dự án thông qua tên đăng nhập hoặc email.",
)
def add_project_member(project_id: int, payload: ProjectMemberAdd, current_user: User = Depends(get_auth_user), db: Session = Depends(get_db)):
    """Thêm thành viên vào dự án bằng username hoặc email."""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Không tìm thấy dự án")

    user_role = _get_user_role(project_id, current_user.id, db)
    if user_role not in ("leader", "deputy") and project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Chỉ Trưởng nhóm hoặc Phó nhóm mới có quyền thêm thành viên vào dự án.")

    # Tìm tài khoản theo username hoặc email
    ident = payload.user_identifier.strip().lower()
    target_user = db.query(User).filter(
        (User.username == ident) | (User.email == ident)
    ).first()

    if not target_user:
        raise HTTPException(status_code=404, detail=f"Không tìm thấy tài khoản người dùng với tên hoặc email '{ident}'.")

    # Kiểm tra xem đã có trong dự án chưa
    existing = db.query(ProjectMember).filter(
        ProjectMember.project_id == project_id,
        ProjectMember.user_id == target_user.id,
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Người dùng '{target_user.full_name}' đã là thành viên trong dự án này.")

    assign_role = payload.role if payload.role in ("deputy", "member") else "member"
    # Chỉ Leader mới được chỉ định trực tiếp làm Deputy khi add
    if assign_role == "deputy" and user_role != "leader":
        assign_role = "member"

    mem = ProjectMember(
        project_id=project_id,
        user_id=target_user.id,
        role=assign_role,
    )
    db.add(mem)
    db.commit()
    db.refresh(mem)

    return {
        "message": f"Đã thêm thành công '{target_user.full_name}' vào dự án với vai trò {ROLE_NAMES.get(assign_role)}.",
        "member_id": mem.id,
        "role": assign_role,
    }


@router.put(
    "/{project_id}/members/{user_id}/role",
    summary="Cập nhật vai trò thành viên (Phân quyền)",
    description="Thay đổi vai trò thành viên (Leader, Deputy, Member). Chỉ DUY NHẤT Trưởng nhóm (Leader) mới có quyền phân quyền.",
)
def update_member_role(project_id: int, user_id: int, payload: ProjectMemberRoleUpdate, current_user: User = Depends(get_auth_user), db: Session = Depends(get_db)):
    """Thay đổi vai trò thành viên trong dự án (Chỉ Leader)."""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Không tìm thấy dự án")

    user_role = _get_user_role(project_id, current_user.id, db)
    if user_role != "leader" and project.owner_id != current_user.id:
        raise HTTPException(
            status_code=403,
            detail="Chỉ duy nhất Trưởng nhóm (Leader) mới có quyền phân quyền hoặc thay đổi vai trò của thành viên trong dự án."
        )

    target_mem = db.query(ProjectMember).filter(
        ProjectMember.project_id == project_id,
        ProjectMember.user_id == user_id,
    ).first()
    if not target_mem:
        raise HTTPException(status_code=404, detail="Thành viên không tồn tại trong dự án này.")

    new_role = payload.role.strip().lower()
    if new_role not in ("leader", "deputy", "member"):
        raise HTTPException(status_code=400, detail="Vai trò không hợp lệ. Chỉ chấp nhận: leader, deputy, member.")

    # Neu trao quyen Leader cho nguoi khac: chuyen nguoi goi thanh deputy hoac giu leader
    if new_role == "leader" and user_id != current_user.id:
        target_mem.role = "leader"
        project.owner_id = user_id
        # Nguoi cu co the tro thanh deputy
        cur_mem = db.query(ProjectMember).filter(
            ProjectMember.project_id == project_id,
            ProjectMember.user_id == current_user.id,
        ).first()
        if cur_mem:
            cur_mem.role = "deputy"
    else:
        target_mem.role = new_role

    db.commit()
    return {
        "message": f"Đã cập nhật vai trò thành viên thành '{ROLE_NAMES.get(target_mem.role)}'.",
        "user_id": user_id,
        "new_role": target_mem.role,
        "new_role_name": ROLE_NAMES.get(target_mem.role),
    }


@router.delete(
    "/{project_id}/members/{user_id}",
    summary="Xóa thành viên khỏi dự án hoặc tự rời dự án",
    description="Leader có thể xóa bất kỳ ai; Deputy có thể xóa Member; Thành viên có thể tự rời khỏi dự án.",
)
def remove_project_member(project_id: int, user_id: int, current_user: User = Depends(get_auth_user), db: Session = Depends(get_db)):
    """Xóa thành viên khỏi dự án."""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Không tìm thấy dự án")

    caller_role = _get_user_role(project_id, current_user.id, db)
    target_mem = db.query(ProjectMember).filter(
        ProjectMember.project_id == project_id,
        ProjectMember.user_id == user_id,
    ).first()

    if not target_mem:
        raise HTTPException(status_code=404, detail="Thành viên không ở trong dự án này.")

    # 1. Tu roi khoi du an
    if current_user.id == user_id:
        if caller_role == "leader":
            # Dem so leader con lai
            other_leaders = db.query(ProjectMember).filter(
                ProjectMember.project_id == project_id,
                ProjectMember.user_id != user_id,
                ProjectMember.role == "leader",
            ).count()
            if other_leaders == 0:
                raise HTTPException(
                    status_code=400,
                    detail="Bạn là Trưởng nhóm duy nhất. Vui lòng chuyển giao vai trò Trưởng nhóm cho thành viên khác trước khi rời dự án hoặc chọn xóa toàn bộ dự án."
                )
        db.delete(target_mem)
        db.commit()
        return {"message": "Bạn đã rời khỏi dự án thành công."}

    # 2. Xoa nguoi khac
    if caller_role == "leader":
        # Leader xoa bat ky ai
        db.delete(target_mem)
        db.commit()
        return {"message": "Đã xóa thành viên khỏi dự án thành công."}
    elif caller_role == "deputy":
        # Deputy chi xoa duoc member
        if target_mem.role in ("leader", "deputy"):
            raise HTTPException(status_code=403, detail="Phó nhóm chỉ có quyền xóa Thành viên thông thường, không thể xóa Trưởng nhóm hoặc Phó nhóm khác.")
        db.delete(target_mem)
        db.commit()
        return {"message": "Đã xóa thành viên khỏi dự án thành công."}
    else:
        raise HTTPException(status_code=403, detail="Bạn không có quyền xóa thành viên khác khỏi dự án.")


# ---------------------------------------------------------------------------
# Xuat Báo Cáo Kiểm Thử Cấp Dự Án (Project-Level Export)
# ---------------------------------------------------------------------------
@router.get(
    "/{project_id}/export",
    summary="Xuất báo cáo kiểm thử toàn diện cấp dự án",
    description="Xuất báo cáo nghiệm thu toàn bộ dự án dưới dạng Excel (4 sheet chuyên nghiệp), CSV, JSON, hoặc Markdown.",
)
def export_project_report(project_id: int, format: str = "xlsx", db: Session = Depends(get_db)):
    """Xuất báo cáo toàn bộ dự án thành file Excel (4 sheet), CSV, JSON, hoặc Markdown."""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Không tìm thấy dự án")

    stats = _calculate_project_stats(project, db)

    # Tap hop toan bo kịch ban va test cases
    breakdown = []
    all_cases = []
    failed_cases = []

    for r in project.requirements:
        r_cases = []
        for s in r.test_suites:
            for tc in s.test_cases:
                item = {
                    "feature": r.title or f"Chức năng #{r.id}",
                    "suite_id": s.id,
                    "code": tc.code,
                    "scenario": tc.scenario,
                    "preconditions": tc.preconditions or "",
                    "test_steps": tc.test_steps or "",
                    "test_type": tc.test_type,
                    "input_data": tc.input_data or {},
                    "expected_result": tc.expected_result,
                    "postconditions": tc.postconditions or "",
                    "actual_result": tc.actual_result or "",
                    "status": tc.status or "Untested",
                    "priority": tc.priority or "Medium",
                    "technique": tc.technique_source or s.technique or "",
                }
                all_cases.append(item)
                r_cases.append(item)
                if tc.status and tc.status.strip().lower() in ("fail", "failed", "không đạt", "lỗi"):
                    failed_cases.append(item)

        pass_c = sum(1 for c in r_cases if c["status"].strip().lower() in ("pass", "passed", "đạt"))
        fail_c = sum(1 for c in r_cases if c["status"].strip().lower() in ("fail", "failed", "không đạt", "lỗi"))
        breakdown.append({
            "id": r.id,
            "title": r.title or f"Chức năng #{r.id}",
            "total_cases": len(r_cases),
            "pass_count": pass_c,
            "fail_count": fail_c,
            "untested_count": len(r_cases) - pass_c - fail_c,
            "pass_rate": round((pass_c / len(r_cases) * 100), 1) if r_cases else 0.0,
        })

    # 1. JSON
    if format == "json":
        data = {
            "project": stats,
            "breakdown": breakdown,
            "test_cases": all_cases,
            "defect_log": failed_cases,
        }
        json_bytes = json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")
        return StreamingResponse(
            io.BytesIO(json_bytes),
            media_type="application/json; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="project_report_{project.code or project.id}.json"'},
        )

    # 2. MARKDOWN
    elif format in ("markdown", "md"):
        lines = [
            f"# Báo Cáo Kiểm Thử Toàn Diện: {project.name}",
            "",
            f"- **Mã Dự Án:** `{project.code}`",
            f"- **Loại Dự Án:** {'Nhóm' if project.project_type == 'team' else 'Cá nhân'}",
            f"- **Phiên Bản:** `{project.version}`",
            f"- **Người Phụ Trách:** {project.lead}",
            f"- **Ngày Báo Cáo:** {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}",
            "",
            "## 1. Tổng Quan Tiến Độ & Đánh Giá Chất Lượng",
            "",
            "| Chỉ Số KPI | Số Lượng / Giá Trị |",
            "|---|---|",
            f"| Tổng số chức năng / yêu cầu | {stats['requirements_count']} |",
            f"| Tổng số ca kiểm thử | {stats['total_cases']} |",
            f"| Số ca Đạt (Pass) | {stats['pass_count']} |",
            f"| Số ca Không Đạt (Fail) | {stats['fail_count']} |",
            f"| Số ca Chưa Kiểm Thử | {stats['untested_count']} |",
            f"| Tỷ Lệ Đạt (Pass Rate) | {stats['pass_rate']}% |",
            f"| Đánh Giá Nghiệm Thu | {'ĐẠT CHUẨN' if stats['pass_rate'] >= 90 else 'CẦN KHẮC PHỤC LỖI'} |",
            "",
            "## 2. Bảng Phân Tích Chi Tiết Theo Chức Năng",
            "",
            "| STT | Tên Chức Năng / Yêu Cầu | Tổng Ca | Đạt (Pass) | Lỗi (Fail) | Chưa Kiểm | Tỷ Lệ Đạt |",
            "|---|---|---|---|---|---|---|",
        ]
        for idx, b in enumerate(breakdown, 1):
            lines.append(f"| {idx} | {b['title']} | {b['total_cases']} | {b['pass_count']} | {b['fail_count']} | {b['untested_count']} | {b['pass_rate']}% |")

        if failed_cases:
            lines.extend([
                "",
                "## 3. Nhật Ký Lỗi Cần Khắc Phục (Defect Log)",
                "",
                "| Mã TC | Chức Năng | Kịch Bản Kiểm Thử | Dữ Liệu Đầu Vào | Kết Quả Mong Đợi | Kết Quả Thực Tế |",
                "|---|---|---|---|---|---|",
            ])
            for fc in failed_cases:
                inp_str = ", ".join(f"`{k}={v}`" for k, v in fc["input_data"].items())
                lines.append(f"| {fc['code']} | {fc['feature']} | {fc['scenario']} | {inp_str} | {fc['expected_result']} | {fc['actual_result']} |")

        md_bytes = "\n".join(lines).encode("utf-8")
        return StreamingResponse(
            io.BytesIO(md_bytes),
            media_type="text/markdown; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="project_report_{project.code or project.id}.md"'},
        )

    # 3. CSV
    elif format == "csv":
        import csv
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "Chức Năng",
            "Mã Test Case",
            "Kịch Bản Kiểm Thử",
            "Tiền Điều Kiện",
            "Các Bước Thực Hiện",
            "Phân Loại",
            "Dữ Liệu Đầu Vào",
            "Kết Quả Mong Đợi",
            "Hậu Điều Kiện",
            "Kết Quả Thực Tế",
            "Trạng Thái",
            "Mức Ưu Tiên",
            "Kỹ Thuật",
        ])
        for c in all_cases:
            inp_str = "\n".join(f"{k} = {v}" for k, v in c["input_data"].items())
            writer.writerow([
                c["feature"],
                c["code"],
                c["scenario"],
                c["preconditions"],
                c["test_steps"],
                c["test_type"].upper(),
                inp_str,
                c["expected_result"],
                c.get("postconditions", ""),
                c["actual_result"],
                c["status"],
                c["priority"],
                c["technique"],
            ])
        output.seek(0)
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="project_report_{project.code or project.id}.csv"'},
        )

    # 4. EXCEL (.xlsx) - 4 SHEETS CHUYÊN NGHIỆP
    else:
        if not HAS_OPENPYXL:
            raise HTTPException(status_code=500, detail="openpyxl not installed")

        wb = openpyxl.Workbook()

        # Styles
        title_font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
        sub_font = Font(name="Segoe UI", size=10, italic=True, color="6B7280")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_fill_accent = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        header_fill_alert = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")

        bold_font = Font(name="Segoe UI", size=10, bold=True)
        cell_font = Font(name="Segoe UI", size=9)
        kpi_num_font = Font(name="Segoe UI", size=18, bold=True, color="1E3A8A")

        thin_side = Side(style="thin", color="D1D5DB")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        card_fill = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")

        # -------------------------------------------------------------
        # SHEET 1: Báo Cáo Tổng Quan (Executive Summary)
        # -------------------------------------------------------------
        ws1 = wb.active
        ws1.title = "Báo Cáo Tổng Quan"

        ws1["A1"] = "BÁO CÁO NGHIỆM THU KIỂM THỬ DỰ ÁN"
        ws1["A1"].font = title_font
        ws1["A2"] = f"Dự án: {project.name} | Mã: {project.code} | Phiên bản: {project.version}"
        ws1["A2"].font = sub_font
        ws1["A3"] = f"Người phụ trách: {project.lead} | Loại: {'Dự Án Nhóm' if project.project_type == 'team' else 'Dự Án Cá Nhân'} | Ngày xuất: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws1["A3"].font = sub_font

        # KPI Blocks
        kpis = [
            ("A5", "B6", "TỔNG CHỨC NĂNG", str(stats["requirements_count"])),
            ("C5", "D6", "TỔNG TEST CASES", str(stats["total_cases"])),
            ("E5", "F6", "SỐ CA ĐẠT (PASS)", str(stats["pass_count"])),
            ("G5", "H6", "SỐ CA LỖI (FAIL)", str(stats["fail_count"])),
            ("I5", "J6", "TỶ LỆ ĐẠT", f"{stats['pass_rate']}%"),
        ]
        for top_l, bot_r, label, val in kpis:
            col_l, row_t = top_l[0], int(top_l[1])
            col_r, row_b = bot_r[0], int(bot_r[1])
            ws1.merge_cells(f"{top_l}:{bot_r}")
            top_cell = ws1[top_l]
            top_cell.value = f"{label}\n{val}"
            top_cell.font = bold_font
            top_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            top_cell.fill = card_fill
            for r in range(row_t, row_b + 1):
                for c in range(ord(col_l) - 64, ord(col_r) - 64 + 1):
                    ws1.cell(row=r, column=c).border = thin_border

        # -------------------------------------------------------------
        # SHEET 2: Ma Trận Chức Năng (Feature Breakdown)
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="Ma Trận Chức Năng")
        headers2 = ["STT", "Mã Chức Năng", "Tên Chức Năng / Yêu Cầu", "Tổng Ca Test", "Đạt (Pass)", "Lỗi (Fail)", "Chưa Kiểm", "Tỷ Lệ Đạt %"]
        for col, h in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill_accent
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for r_idx, b in enumerate(breakdown, 2):
            vals = [r_idx - 1, f"REQ-{b['id']:03d}", b["title"], b["total_cases"], b["pass_count"], b["fail_count"], b["untested_count"], f"{b['pass_rate']}%"]
            for col, v in enumerate(vals, 1):
                cell = ws2.cell(row=r_idx, column=col, value=v)
                cell.font = cell_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center" if col in (1, 2, 4, 5, 6, 7, 8) else "left", vertical="center")

        # -------------------------------------------------------------
        # SHEET 3: Toàn Bộ Test Cases (All Test Cases)
        # -------------------------------------------------------------
        ws3 = wb.create_sheet(title="Toàn Bộ Test Cases")
        headers3 = ["Chức Năng", "Mã TC", "Kịch Bản Kiểm Thử", "Tiền Điều Kiện", "Các Bước", "Phân Loại", "Dữ Liệu Đầu Vào", "Kết Quả Mong Đợi", "Hậu Điều Kiện", "Kết Quả Thực Tế", "Trạng Thái", "Mức Ưu Tiên", "Kỹ Thuật"]
        for col, h in enumerate(headers3, 1):
            cell = ws3.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for r_idx, tc in enumerate(all_cases, 2):
            inp_s = "\n".join(f"{k} = {v}" for k, v in tc["input_data"].items())
            vals = [tc["feature"], tc["code"], tc["scenario"], tc["preconditions"], tc["test_steps"], tc["test_type"].upper(), inp_s, tc["expected_result"], tc.get("postconditions", ""), tc["actual_result"], tc["status"], tc["priority"], tc["technique"]]
            for col, v in enumerate(vals, 1):
                cell = ws3.cell(row=r_idx, column=col, value=v)
                cell.font = cell_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center" if col in (2, 6, 10, 11, 12) else "left", vertical="center", wrap_text=True)

        # -------------------------------------------------------------
        # SHEET 4: Nhật Ký Lỗi (Defect Log)
        # -------------------------------------------------------------
        ws4 = wb.create_sheet(title="Nhật Ký Lỗi (Defects)")
        headers4 = ["STT", "Chức Năng", "Mã TC", "Kịch Bản Lỗi", "Dữ Liệu Đầu Vào", "Kết Quả Mong Đợi", "Kết Quả Thực Tế Gặp Lỗi", "Mức Ưu Tiên"]
        for col, h in enumerate(headers4, 1):
            cell = ws4.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill_alert
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for r_idx, fc in enumerate(failed_cases, 2):
            inp_s = "\n".join(f"{k} = {v}" for k, v in fc["input_data"].items())
            vals = [r_idx - 1, fc["feature"], fc["code"], fc["scenario"], inp_s, fc["expected_result"], fc["actual_result"], fc["priority"]]
            for col, v in enumerate(vals, 1):
                cell = ws4.cell(row=r_idx, column=col, value=v)
                cell.font = cell_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center" if col in (1, 3, 8) else "left", vertical="center", wrap_text=True)

        # Set column widths for each sheet
        for ws in [ws1, ws2, ws3, ws4]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="project_report_{project.code or project.id}.xlsx"'},
        )


@router.post(
    "/seed-sample",
    response_model=ProjectOut,
    summary="Nạp dự án mẫu E-Commerce Scrum để nghiệm thu đồ án",
    description="Khởi tạo tự động một dự án mẫu Scrum hoàn chỉnh gồm 3 chức năng con, Sprint, phân công thành viên và 141 ca kiểm thử thực tế phục vụ demo nghiệm thu đồ án CDIO-4.",
)
def seed_sample_project(current_user: User = Depends(get_auth_user), db: Session = Depends(get_db)):
    """Tạo một dự án mẫu hoàn chỉnh phục vụ báo cáo / demo đồ án CDIO-4."""
    existing = db.query(Project).filter(Project.code == "PRJ-SHOP-DEMO").first()
    if existing:
        # Dam bao current_user co trong project
        mem = db.query(ProjectMember).filter(ProjectMember.project_id == existing.id, ProjectMember.user_id == current_user.id).first()
        if not mem:
            mem = ProjectMember(project_id=existing.id, user_id=current_user.id, role="leader")
            db.add(mem)
            db.commit()
        return _calculate_project_stats(existing, db, current_user.id)

    import datetime
    today = datetime.date.today()

    # 1. Tao Du An Nhom
    p = Project(
        name="Hệ Thống Đặt Hàng & Thanh Toán E-Commerce (Dự Án Mẫu)",
        code="PRJ-SHOP-DEMO",
        description="Dự án mẫu Scrum phục vụ demo nghiệm thu đồ án CDIO-4. Bao gồm đầy đủ Sprint, User Stories, phân công thành viên và bộ Test Case hoàn chỉnh.",
        version="1.2.0",
        lead=current_user.full_name,
        owner_id=current_user.id,
        project_type="team",
        join_code="TEAM-DEMO88",
    )
    db.add(p)
    db.flush()

    # Them Leader cho current_user
    db.add(ProjectMember(project_id=p.id, user_id=current_user.id, role="leader"))
    # Them thanh vien mau khac neu co trong DB
    other_users = db.query(User).filter(User.id != current_user.id).all()
    for idx, u in enumerate(other_users):
        r = "deputy" if idx == 0 else "member"
        db.add(ProjectMember(project_id=p.id, user_id=u.id, role=r))

    db.flush()

    # 2. Tao Sprint 1
    from models import Sprint
    sprint1 = Sprint(
        project_id=p.id,
        name="Sprint 1: Xác Thực & Mua Sắm E-Commerce",
        goal="Nghiệm thu toàn diện luồng Đăng ký tài khoản, Áp mã khuyến mãi và Cổng thanh toán trực tuyến",
        start_date=today.strftime("%Y-%m-%d"),
        end_date=(today + datetime.timedelta(days=14)).strftime("%Y-%m-%d"),
        status="active",
    )
    db.add(sprint1)
    db.flush()

    # 3. Tao 3 Chuc Nang Con (Requirements)
    features_data = [
        {
            "title": "Xác thực tài khoản & Phân quyền vai trò",
            "raw_text": (
                "Khi người dùng đăng ký tài khoản:\n"
                "1. Tuổi phải từ 18 đến 60 tuổi.\n"
                "2. Mật khẩu phải từ 8 đến 20 ký tự, có ít nhất 1 chữ hoa và 1 số.\n"
                "3. Vai trò người dùng thuộc một trong các giá trị: Admin, Member, Guest."
            ),
            "assignee": "Nguyễn Văn A (Tester)",
            "due_date": (today + datetime.timedelta(days=5)).strftime("%Y-%m-%d"),
            "priority": "High",
            "status": "completed",
            "params": [
                {"name": "tuoi", "data_type": "integer", "min_value": 18, "max_value": 60, "is_required": True, "description": "Tuổi người dùng"},
                {"name": "mat_khau", "data_type": "string", "min_length": 8, "max_length": 20, "is_required": True, "description": "Mật khẩu"},
                {"name": "vai_tro", "data_type": "enum", "enum_values": ["Admin", "Member", "Guest"], "is_required": True, "description": "Vai trò"},
            ]
        },
        {
            "title": "Quản lý giỏ hàng & Áp mã Voucher khuyến mãi",
            "raw_text": (
                "Khi người dùng áp mã giảm giá trong giỏ hàng:\n"
                "1. Số lượng sản phẩm từ 1 đến 50 cái.\n"
                "2. Tổng giá trị đơn hàng từ 100000 đến 15000000 VNĐ.\n"
                "3. Tỷ lệ giảm giá voucher từ 5 đến 40%.\n"
                "4. Hình thức nhận hàng: Giao hàng tiêu chuẩn, Hỏa tốc 2h."
            ),
            "assignee": "Trần Minh (QA Lead)",
            "due_date": (today + datetime.timedelta(days=8)).strftime("%Y-%m-%d"),
            "priority": "High",
            "status": "completed",
            "params": [
                {"name": "so_luong", "data_type": "integer", "min_value": 1, "max_value": 50, "is_required": True, "description": "Số lượng sản phẩm"},
                {"name": "tong_tien", "data_type": "float", "min_value": 100000, "max_value": 15000000, "is_required": True, "description": "Tổng giá trị đơn hàng"},
                {"name": "voucher_pct", "data_type": "integer", "min_value": 5, "max_value": 40, "is_required": True, "description": "Tỷ lệ giảm giá"},
                {"name": "hinh_thuc_giao", "data_type": "enum", "enum_values": ["Giao hàng tiêu chuẩn", "Hỏa tốc 2h"], "is_required": True, "description": "Hình thức giao"},
            ]
        },
        {
            "title": "Cổng thanh toán trực tuyến & Quy tắc giao dịch",
            "raw_text": (
                "Khi khách hàng tiến hành thanh toán:\n"
                "1. Số tiền thanh toán từ 50000 đến 50000000 VNĐ.\n"
                "2. Phương thức thanh toán: Ví Momo, Chuyển khoản QR, Thẻ tín dụng Visa.\n"
                "3. Số lần nhập sai mã OTP tối đa từ 1 đến 3 lần."
            ),
            "assignee": "Lê Hoàng (Tester)",
            "due_date": (today + datetime.timedelta(days=12)).strftime("%Y-%m-%d"),
            "priority": "Medium",
            "status": "in_testing",
            "params": [
                {"name": "so_tien", "data_type": "float", "min_value": 50000, "max_value": 50000000, "is_required": True, "description": "Số tiền thanh toán"},
                {"name": "phuong_thuc", "data_type": "enum", "enum_values": ["Ví Momo", "Chuyển khoản QR", "Thẻ tín dụng Visa"], "is_required": True, "description": "Phương thức"},
                {"name": "so_lan_sai_otp", "data_type": "integer", "min_value": 1, "max_value": 3, "is_required": True, "description": "Số lần sai OTP"},
            ]
        }
    ]

    from engine.assembler import assemble_test_suite
    from models import Parameter, TestSuite, TestCase

    for feat in features_data:
        req = Requirement(
            project_id=p.id,
            sprint_id=sprint1.id,
            title=feat["title"],
            raw_text=feat["raw_text"],
            format_type="free_text",
            assignee_name=feat["assignee"],
            due_date=feat["due_date"],
            priority=feat["priority"],
            status=feat["status"],
        )
        db.add(req)
        db.flush()

        for p_data in feat["params"]:
            param = Parameter(
                requirement_id=req.id,
                name=p_data["name"],
                data_type=p_data["data_type"],
                min_value=p_data.get("min_value"),
                max_value=p_data.get("max_value"),
                min_length=p_data.get("min_length"),
                max_length=p_data.get("max_length"),
                enum_values=p_data.get("enum_values"),
                is_required=p_data.get("is_required", True),
                description=p_data.get("description", ""),
            )
            db.add(param)

        db.flush()

        # Sinh bo test suite tu dong
        result = assemble_test_suite(
            parameters=feat["params"],
            techniques=["bva", "equivalence_partitioning", "pairwise"],
            requirement_title=req.title,
        )

        test_suite = TestSuite(
            requirement_id=req.id,
            technique="BVA + EP + Pairwise",
            total_cases=len(result["test_cases"]),
            coverage_info=result["coverage_info"],
        )
        db.add(test_suite)
        db.flush()

        # Luu cac test cases va tao ket qua mau Pass/Fail thuc te (khoang 92% Pass, 8% Fail)
        for idx, tc_data in enumerate(result["test_cases"]):
            is_fail = (idx % 12 == 11)  # Ca thu 12 thi bi fail de tao su thuc te cho defect log
            status_val = "Fail" if is_fail else "Pass"
            actual_val = "Lỗi hệ thống không xử lý đúng biên trên" if is_fail else "Khớp hoàn toàn với kết quả mong đợi"

            tc = TestCase(
                test_suite_id=test_suite.id,
                code=tc_data["code"],
                scenario=tc_data["scenario"],
                preconditions=tc_data.get("preconditions", ""),
                test_steps=tc_data.get("test_steps", ""),
                test_type=tc_data["test_type"],
                input_data=tc_data["input_data"],
                expected_result=tc_data["expected_result"],
                postconditions=tc_data.get("postconditions", ""),
                actual_result=actual_val,
                status=status_val,
                priority=tc_data.get("priority", "Medium"),
                technique_source=tc_data.get("technique_source", ""),
            )
            db.add(tc)

    db.commit()
    db.refresh(p)
    return _calculate_project_stats(p, db, current_user.id)

